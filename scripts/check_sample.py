"""Recheck five fixed app cells against the official source workbook."""

import argparse
from contextlib import closing
from pathlib import Path
import sqlite3

from openpyxl import load_workbook


CHECKS = (
    (2, "men", "PARTIC_MEN", "participants_raw"),
    (2, "men", "OPEXPPERTEAM_MEN", "team_expenses_raw"),
    (3, "women", "PARTIC_WOMEN", "participants_raw"),
    (3, "women", "OPEXPPERTEAM_WOMEN", "team_expenses_raw"),
    (4, "men", "PARTIC_MEN", "participants_raw"),
)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", type=Path, default=Path("data/raw/2024/schools.xlsx"))
    parser.add_argument("--db", type=Path, default=Path("data/processed/eada.sqlite"))
    args = parser.parse_args()
    workbook = load_workbook(args.input, read_only=True, data_only=True)
    try:
        sheet = workbook["schools"]
        headers = [cell.value for cell in sheet[1]]
        indexes = {name: index for index, name in enumerate(headers)}
        with closing(sqlite3.connect(args.db)) as connection:
            for row_number, category, column, field in CHECKS:
                source = next(sheet.iter_rows(min_row=row_number, max_row=row_number, values_only=True))
                expected = str(source[indexes[column]])
                stored = connection.execute(
                    f"SELECT {field} FROM eada WHERE source_row = ? AND category = ?",
                    (row_number, category),
                ).fetchone()
                if stored is None or stored[0] != expected:
                    raise AssertionError(f"row {row_number} {column}: source={expected}, app={stored}")
                print(f"PASS row {row_number} {column}: {expected}")
    finally:
        workbook.close()


if __name__ == "__main__":
    main()
