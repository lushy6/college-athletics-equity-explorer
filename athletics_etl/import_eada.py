"""Read one EADA schools file without losing source cell meaning."""

import csv
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any, Iterator

from openpyxl import load_workbook

from .schema import ImportResult, NumericCell, Record, RejectedRow


class SchemaError(ValueError):
    """The selected file cannot be mapped safely to the known EADA layout."""


REQUIRED = (
    "unitid", "institution_name", "state_cd", "classification_name",
    "sector_name", "SPORTSCODE", "Sports", "PARTIC_MEN", "PARTIC_WOMEN",
    "OPEXPPERPART_MEN", "OPEXPPERTEAM_MEN", "OPEXPPERPART_WOMEN",
    "OPEXPPERTEAM_WOMEN",
)


def _source_rows(path: Path) -> Iterator[tuple[int, dict[str, Any]]]:
    if path.suffix.lower() == ".csv":
        with path.open(newline="", encoding="utf-8-sig") as source:
            reader = csv.DictReader(source)
            columns = reader.fieldnames or []
            _check_columns(columns)
            for number, row in enumerate(reader, start=2):
                yield number, row
    elif path.suffix.lower() == ".xlsx":
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            sheet = workbook["schools"] if "schools" in workbook else workbook.active
            rows = sheet.iter_rows(values_only=True)
            columns = [str(value).strip() if value is not None else "" for value in next(rows)]
            _check_columns(columns)
            for number, values in enumerate(rows, start=2):
                yield number, dict(zip(columns, values))
        except StopIteration as exc:
            raise SchemaError("empty workbook") from exc
        finally:
            workbook.close()
    else:
        raise SchemaError(f"unsupported file type: {path.suffix}")


def _check_columns(columns: list[str]) -> None:
    missing = sorted(set(REQUIRED) - set(columns))
    if missing:
        raise SchemaError("missing columns: " + ", ".join(missing))
    if len(columns) != len(set(columns)):
        raise SchemaError("duplicate column names")


def _raw(value: Any) -> str:
    return "" if value is None else str(value).strip()


def _numeric(value: Any, *, count: bool = False) -> NumericCell:
    raw = _raw(value)
    if raw == "":
        return NumericCell(raw, None, "blank")
    if raw.upper() in {"D", "S", "SUPPRESSED"}:
        return NumericCell(raw, None, "suppressed")
    if raw.upper() in {"N/A", "NA", "NULL", "UNDEFINED"}:
        return NumericCell(raw, None, "undefined")
    try:
        parsed = Decimal(raw.replace(",", ""))
    except InvalidOperation:
        return NumericCell(raw, None, "invalid")
    if not parsed.is_finite() or parsed < 0 or (count and parsed != parsed.to_integral_value()):
        return NumericCell(raw, None, "invalid")
    return NumericCell(raw, parsed, "observed")


def _institution_id(value: Any) -> str | None:
    raw = _raw(value)
    if raw.endswith(".0"):
        raw = raw[:-2]
    return raw if raw.isascii() and raw.isdigit() and len(raw) == 6 else None


def import_year(path: str | Path, year: int) -> ImportResult:
    """Normalize one collection year; reject bad/duplicate keys explicitly."""
    source = Path(path)
    if not 2000 <= year <= 2100:
        raise ValueError("year must be a collection year")
    records: list[Record] = []
    rejects: list[RejectedRow] = []
    seen: set[tuple[int, str, str]] = set()
    source_rows = 0
    for row_number, row in _source_rows(source):
        source_rows += 1
        institution_id = _institution_id(row["unitid"])
        if institution_id is None:
            rejects.append(RejectedRow(source.name, row_number, "invalid institution ID"))
            continue
        sport_code = _raw(row["SPORTSCODE"])
        if sport_code.endswith(".0"):
            sport_code = sport_code[:-2]
        if not sport_code.isascii() or not sport_code.isdigit():
            rejects.append(RejectedRow(source.name, row_number, "invalid sport code"))
            continue
        key = (year, institution_id, sport_code)
        if key in seen:
            rejects.append(RejectedRow(source.name, row_number, "duplicate institution/sport/year"))
            continue
        seen.add(key)
        for category, suffix in (("men", "MEN"), ("women", "WOMEN")):
            records.append(Record(
                year=year, institution_id=institution_id,
                institution_name=_raw(row["institution_name"]),
                state=_raw(row["state_cd"]), division=_raw(row["classification_name"]),
                institution_type=_raw(row["sector_name"]), sport_code=sport_code,
                sport_name=_raw(row["Sports"]), category=category,
                participants=_numeric(row[f"PARTIC_{suffix}"], count=True),
                team_expenses=_numeric(row[f"OPEXPPERTEAM_{suffix}"]),
                reported_per_participant=_numeric(row[f"OPEXPPERPART_{suffix}"]),
                source_file=source.name, source_row=row_number,
            ))
    return ImportResult(records, rejects, source_rows)
